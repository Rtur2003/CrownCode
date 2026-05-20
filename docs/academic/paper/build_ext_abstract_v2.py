"""
Genişletilmiş İngilizce Özet — TEK SAYFA, tüm hücreler dolu, Graphical Abstract figürü ile.
Şablonun karmaşık merge yapısını bırak; sade yeni bir doküman oluştur.
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FIGURES = HERE.parent / "figures"
OUT = HERE / "deliverables" / "AURIS_Genisletilmis_Ingilizce_Ozet.docx"


def _font(run, *, size=9, bold=False, italic=False, color="333333", name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def _clear(cell):
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    cell.add_paragraph()


def cell_text(cell, text, *, size=9, bold=False, italic=False, align=None):
    _clear(cell)
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    r = p.add_run(text)
    _font(r, size=size, bold=bold, italic=italic)


def cell_multi(cell, lines, *, size=9, bold=False, italic=False):
    _clear(cell)
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    r = p.add_run(lines[0])
    _font(r, size=size, bold=bold, italic=italic)
    for line in lines[1:]:
        np = cell.add_paragraph()
        np.paragraph_format.space_after = Pt(0)
        np.paragraph_format.line_spacing = 1.15
        r = np.add_run(line)
        _font(r, size=size, bold=bold, italic=italic)


def cell_image(cell, img_path, width_cm=6.5):
    _clear(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    if img_path.exists():
        r = p.add_run()
        r.add_picture(str(img_path), width=Cm(width_cm))


doc = Document()

# A4 + 2.5 cm marjinler
s = doc.sections[0]
s.page_width = Cm(21.0)
s.page_height = Cm(29.7)
s.top_margin = Cm(2.5)
s.bottom_margin = Cm(2.5)
s.left_margin = Cm(2.5)
s.right_margin = Cm(2.5)

# ─── Başlık (EN üstte, TR altta italik) ───
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("AURIS: A Multi-Model Ensemble Approach for the Detection of "
              "AI-Generated Music")
_font(r, size=14, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("AURIS: Çoklu-Model Topluluk Yaklaşımı ile Yapay Zekâ Tarafından "
              "Üretilen Müziklerin Tespiti")
_font(r, size=11, italic=True)

# ─── Yazar bloğu ───
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Hasan Arthur Altuntaş1,*")
_font(r, size=10, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run("1Department of Computer Engineering, Faculty of Engineering, "
              "Düzce University, 81620, Düzce, Türkiye")
_font(r, size=9, italic=True)

# ─── İki kolonlu ana tablo: SOL (Highlights, Keywords, Article Info, Ack, Correspondence)
#                              SAĞ (Graphical Abstract + Extended Abstract bölümleri)
tbl = doc.add_table(rows=1, cols=2)
tbl.style = "Table Grid"
tbl.autofit = False

# Sütun genişlikleri
left_w, right_w = Cm(6.0), Cm(10.0)
for cell in tbl.rows[0].cells:
    pass
tbl.rows[0].cells[0].width = left_w
tbl.rows[0].cells[1].width = right_w

# ─── SOL HÜCRE ───
left = tbl.rows[0].cells[0]
_clear(left)

# Highlights başlığı
p = left.paragraphs[0]
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Highlights:")
_font(r, size=9, bold=True)

# Highlights bullets
for h in [
    "• End-to-end AI music detection with 47-D acoustic features and 11-model "
    "ensemble on 5,195 samples.",
    "• LightGBM achieves 95.48% ROC-AUC with the lowest cross-fold variance "
    "(±0.0023).",
    "• Spectral flatness standard deviation is the most informative feature "
    "for AI-vs-human distinction.",
]:
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.2)
    r = p.add_run(h)
    _font(r, size=9)

# Keywords başlığı
p = left.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Keywords:")
_font(r, size=9, bold=True)

for kw in ["AI-generated music", "Deep learning", "Gradient boosting",
           "Ensemble learning", "Spectral flatness"]:
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.2)
    r = p.add_run(kw)
    _font(r, size=9)

# Article Info başlığı
p = left.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Article Info:")
_font(r, size=9, bold=True)

for line in ["Research Article", "Received: dd.mm.yyyy", "Accepted: dd.mm.yyyy",
             "DOI:"]:
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.2)
    r = p.add_run(line)
    _font(r, size=9)

# Acknowledgement
p = left.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Acknowledgement:")
_font(r, size=9, bold=True)

p = left.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.15
p.paragraph_format.left_indent = Cm(0.2)
r = p.add_run("The author received no specific funding for this work.")
_font(r, size=9, italic=True)

# Correspondence
p = left.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Correspondence:")
_font(r, size=9, bold=True)

for line in ["Author: Hasan Arthur Altuntaş", "e-mail: hasannarthurrr@gmail.com",
             "ORCID: 0009-0002-8302-7657"]:
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.2)
    r = p.add_run(line)
    _font(r, size=9)

# ─── SAĞ HÜCRE — Graphical Abstract + Extended Abstract bölümleri ───
right = tbl.rows[0].cells[1]
_clear(right)

# Graphical/Tabular Abstract başlığı
p = right.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Graphical/Tabular Abstract")
_font(r, size=10, bold=True)

# Görsel
p = right.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
img = FIGURES / "paper_fold_std_table.png"
if img.exists():
    r = p.add_run()
    r.add_picture(str(img), width=Cm(9.5))

# Caption
p = right.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Figure A. Per-fold ROC-AUC of the eleven models, LightGBM at the top.")
_font(r, size=8, italic=True)

# Purpose
p = right.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Purpose: ")
_font(r, size=9, bold=True)
r = p.add_run(
    "The widespread diffusion of text-to-music generators such as Suno, Udio "
    "and MusicGen has created a detection problem of practical importance "
    "for copyright attribution and streaming-platform integrity. This study "
    "aims to design an end-to-end detection system that distinguishes "
    "AI-generated music from human composition with competitive accuracy "
    "and operational interpretability.")
_font(r, size=9)

# Theory and Methods
p = right.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Theory and Methods: ")
_font(r, size=9, bold=True)
r = p.add_run(
    "A 47-dimensional acoustic feature vector covering spectral, temporal, "
    "harmonic-tonal, MFCC and vocal families is extracted with librosa. An "
    "ensemble of eleven classifiers (Logistic Regression, Random Forest, "
    "Gradient Boosting, SVM-RBF, MLP, XGBoost, LightGBM, Deep MLP, 1D-CNN, "
    "Residual MLP, Attention MLP) is trained on 5,195 samples from twelve AI "
    "generators (Suno, Udio, MusicGen, AudioLDM2, Stable Audio, Riffusion, "
    "Mustango, JEN-1, and others) and human sources (GTZAN, FMA, "
    "SleepyJesse covers) under stratified 5-fold cross-validation. The "
    "Youden-J statistic optimises the decision threshold; Brier score and "
    "SHAP analysis are used for calibration and interpretability.")
_font(r, size=9)

# Results
p = right.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Results: ")
_font(r, size=9, bold=True)
r = p.add_run(
    "LightGBM achieves the highest mean ROC-AUC at 0.9548 with the lowest "
    "fold-to-fold standard deviation (±0.0023); its five folds span only "
    "0.9515-0.9580. Deep MLP follows narrowly at 0.9542. Spectral-flatness "
    "standard deviation ranks first in feature importance. The "
    "Youden-optimal threshold θ* = 0.4316 yields a balanced confusion "
    "matrix (TN=2721, FP=392, FN=220, TP=1862) with 89.4% AI sensitivity "
    "and 87.4% human specificity. A Brier score of 0.083 confirms "
    "well-calibrated probabilities.")
_font(r, size=9)

# Conclusion
p = right.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Conclusion: ")
_font(r, size=9, bold=True)
r = p.add_run(
    "AURIS demonstrates that a handcrafted 47-D feature vector with a "
    "heterogeneous ensemble is competitive with deep-learning alternatives "
    "for AI music detection. Two limitations stand out: tree-based models "
    "exhibit an 8-14 percentage-point train-CV accuracy gap, and the bottom "
    "17 features add no measurable accuracy. Future work will pursue "
    "cross-generator evaluation on SONICS and FakeMusicCaps, wav2vec2 "
    "integration, adversarial-robustness testing, and dataset expansion.")
_font(r, size=9)

# Kaydet
doc.save(str(OUT))
print(f"Saved: {OUT.name}")
