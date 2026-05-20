"""Genişletilmiş İngilizce Özet — şablonun tüm hücrelerini doldurur."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent
SRC = Path(r"D:/Downloads/Genişletilmiş İngilizce Özet.docx")
OUT = HERE / "deliverables" / "AURIS_Genisletilmis_Ingilizce_Ozet.docx"
FIGURES = HERE.parent / "figures"
OUT.parent.mkdir(parents=True, exist_ok=True)


def _set_font(run, *, size=9, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def _clear_cell(cell):
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    cell.add_paragraph()


def fill(cell, text, *, size=9, bold=False, italic=False, align=None):
    _clear_cell(cell)
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    _set_font(r, size=size, bold=bold, italic=italic)


def fill_multi(cell, lines, *, size=9, bold=False):
    _clear_cell(cell)
    p = cell.paragraphs[0]
    r = p.add_run(lines[0])
    _set_font(r, size=size, bold=bold)
    for line in lines[1:]:
        p2 = cell.add_paragraph()
        r2 = p2.add_run(line)
        _set_font(r2, size=size, bold=bold)


def fill_image(cell, img_path, width_cm=6.0):
    _clear_cell(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if img_path.exists():
        r = p.add_run()
        r.add_picture(str(img_path), width=Cm(width_cm))


doc = Document(str(SRC))

# Tablo 0: Başlık
tbl0 = doc.tables[0]
fill(tbl0.rows[1].cells[0],
     "AURIS: A Multi-Model Ensemble Approach for the Detection of AI-Generated Music",
     size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
fill(tbl0.rows[2].cells[0],
     "AURIS: Çoklu-Model Topluluk Yaklaşımı ile Yapay Zekâ Tarafından Üretilen "
     "Müziklerin Tespiti",
     size=14, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

# Tablo 1
tbl = doc.tables[1]

# R0: Yazar (her 3 sütun)
for c in range(3):
    fill(tbl.rows[0].cells[c], "Hasan Arthur Altuntaş1,*",
         size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

# R1: Kurum (her 3 sütun)
for c in range(3):
    fill(tbl.rows[1].cells[c],
         "1Department of Computer Engineering, Faculty of Engineering, "
         "Düzce University, 81620, Düzce, Türkiye",
         size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

# R3: Highlights başlığı + Graphical/Tabular Abstract
fill(tbl.rows[3].cells[0], "Highlights:", size=9, bold=True)
fill(tbl.rows[3].cells[2], "Graphical/Tabular Abstract", size=9, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER)

# R4: Highlights bullets + Graphical Abstract (görsel — 8 numaralı fold std table güzel olur)
fill_multi(tbl.rows[4].cells[0], [
    "• An end-to-end system distinguishing AI-generated music from human composition "
    "via a 47-dimensional acoustic feature vector and 5,195 audio samples.",
    "• Seven ML and four DL algorithms compared under a unified 5-fold "
    "cross-validation protocol; LightGBM achieves 95.48% ROC-AUC with the lowest "
    "cross-fold variance (±0.0023).",
    "• Spectral flatness standard deviation emerges as the single most informative "
    "feature; an 8-14 point train-CV accuracy gap reveals overfitting in tree-based "
    "models, and roughly 17 of the 47 features contribute no measurable accuracy.",
], size=9)
fill_image(tbl.rows[4].cells[2], FIGURES / "paper_fold_std_table.png", width_cm=6.5)

# R5: Keywords başlığı + Purpose
fill(tbl.rows[5].cells[0], "Keywords:", size=9, bold=True)
fill(tbl.rows[5].cells[2], "Purpose:", size=9, bold=True)

# R6: Keywords + Purpose text
fill_multi(tbl.rows[6].cells[0], [
    "AI-generated music",
    "Deep learning",
    "Gradient boosting",
    "Ensemble learning",
    "Spectral flatness",
], size=9)
fill(tbl.rows[6].cells[2],
     "The widespread diffusion of text-to-music generators such as Suno, Udio and "
     "MusicGen has created a detection problem of practical importance for copyright "
     "attribution, streaming-platform integrity and artist economics. The present "
     "study aims to design an end-to-end system that distinguishes AI-generated "
     "music from human composition with both competitive accuracy and operational "
     "interpretability.",
     size=9)

# R7: Article Info + Theory and Methods
fill(tbl.rows[7].cells[0], "Article Info:", size=9, bold=True)
fill(tbl.rows[7].cells[2], "Theory and Methods:", size=9, bold=True)

# R8: Received/DOI + theory text
fill_multi(tbl.rows[8].cells[0], [
    "Research Article",
    "Received: dd.mm.yyyy",
    "Accepted: dd.mm.yyyy",
    "",
    "DOI:",
], size=9)
fill(tbl.rows[8].cells[2],
     "A 47-dimensional acoustic feature vector covering spectral, temporal, "
     "harmonic-tonal, MFCC and vocal families is extracted using librosa. An "
     "ensemble of eleven classifiers — Logistic Regression, Random Forest, Gradient "
     "Boosting, SVM-RBF, MLP Neural Network, XGBoost, LightGBM, Deep MLP, 1D-CNN, "
     "Residual MLP and Attention MLP — is trained on 5,195 samples drawn from "
     "twelve AI generation systems (Suno, Udio, MusicGen, AudioLDM2, Stable Audio, "
     "Riffusion, Mustango, JEN-1, and others) and multiple human sources (GTZAN, "
     "FMA, SleepyJesse covers) under a stratified 5-fold cross-validation protocol. "
     "The Youden-J statistic is used to optimise the decision threshold; Brier "
     "score and SHAP analysis are employed for calibration and interpretability "
     "evaluation respectively.",
     size=9)

# R9: boş + Results
fill(tbl.rows[9].cells[0], "", size=9)
fill(tbl.rows[9].cells[2], "Results:", size=9, bold=True)

# R10: boş + results text
fill(tbl.rows[10].cells[0], "", size=9)
fill(tbl.rows[10].cells[2],
     "LightGBM achieves the highest mean ROC-AUC at 0.9548 with the lowest "
     "fold-to-fold standard deviation (±0.0023) in the entire pool; its five folds "
     "span only 0.9515-0.9580. Deep MLP follows narrowly at 0.9542. The ML "
     "family (mean AUC 0.9275) and the DL family (0.9232) are essentially "
     "equivalent when the 1D-CNN architectural mismatch is excluded. "
     "Spectral-flatness standard deviation ranks first in feature importance, "
     "consistent with the systematic difference between synthetic and recorded "
     "spectra. Youden-optimal threshold θ* = 0.4316 yields a balanced confusion "
     "matrix (TN=2721, FP=392, FN=220, TP=1862) with 89.4% AI sensitivity and "
     "87.4% human specificity. A Brier score of 0.083 confirms well-calibrated "
     "probabilities. Per-source analysis reveals 93.0% recall on Suno and 88.6% on "
     "Echoes, but only 50.0% on the deepfake subset, exposing the cross-generator "
     "generalisation challenge.",
     size=9)

# R11: Acknowledgement + Conclusion başlığı
fill(tbl.rows[11].cells[0], "Acknowledgement:", size=9, bold=True)
fill(tbl.rows[11].cells[2], "Conclusion:", size=9, bold=True)

# R12: boş + conclusion text
fill(tbl.rows[12].cells[0], "", size=9)
fill(tbl.rows[12].cells[2],
     "The proposed AURIS system demonstrates that a handcrafted 47-dimensional "
     "feature vector combined with a heterogeneous ensemble of classifiers is "
     "competitive with deep-learning alternatives for AI-generated music "
     "detection. Diagnostic analysis exposes two key limitations: tree-based "
     "models exhibit an 8-14 percentage-point train-CV accuracy gap, and the "
     "bottom 17 features add no measurable accuracy, indicating clear "
     "opportunities for feature pruning and overfitting mitigation. Future work "
     "will pursue formal cross-generator held-out evaluation on emerging "
     "benchmarks (SONICS, FakeMusicCaps), integration of fine-tuned wav2vec2 "
     "into the protocol, explicit adversarial-robustness testing under MP3 "
     "compression and pitch shifting, and expansion of the dataset toward ten "
     "thousand samples covering emerging generation systems.",
     size=9)

# R13: boş
fill(tbl.rows[13].cells[0], "", size=9)
fill(tbl.rows[13].cells[2], "", size=9)

# R14: Correspondence
fill(tbl.rows[14].cells[0], "Correspondence:", size=9, bold=True)
fill(tbl.rows[14].cells[2], "", size=9)

# R15: Author/email
fill_multi(tbl.rows[15].cells[0], [
    "Author: Hasan Arthur Altuntaş",
    "e-mail: hasannarthurrr@gmail.com",
    "ORCID: 0009-0002-8302-7657",
], size=9)
fill(tbl.rows[15].cells[2], "", size=9)

# R16: boş
fill(tbl.rows[16].cells[0], "", size=9)
fill(tbl.rows[16].cells[2], "", size=9)

# R17: YAZARA NOT — şablonda zaten var, dokunma

doc.save(str(OUT))
print(f"Saved: {OUT}")
