"""Kapak Sayfası — Resmi şablon (Kapak Sayfası.docx) doldurulur."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(r"D:/Downloads/Kapak Sayfası.docx")
OUT = Path(__file__).resolve().parent / "deliverables" / "AURIS_Kapak_Sayfasi.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


def _set_font(run, *, size=9, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def replace_para(p, text, *, size=9, bold=False):
    """Replace all runs in paragraph with single new run."""
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(text)
    _set_font(r, size=size, bold=bold)


doc = Document(str(SRC))

# Şablon yapısı (incelemeden bildiğimiz indeksler):
# P0: "Türkçe makale başlığı:"  (etiket - dokunma)
# P1: TR başlık (değiştirilecek)
# P3: "Yazar adları ve adres bilgileri:"  (etiket)
# P4: Yazar adı
# P5: Kurum
# P6: 2. kurum (silinecek - tek kurumumuz var)
# P7: ORCID
# P8: E-posta
# P9: "İletişim yazarı telefon no:"
# P12: "İngilizce makale başlığı:"  (etiket)
# P13: EN başlık
# P16: "Yazar adları ve adres bilgileri:"  (etiket)
# P17: Yazar
# P18: Kurum EN
# P20: ORCID
# P21: E-posta
# P22: "Phone of contact author:"

paragraphs = doc.paragraphs

# TR başlık
replace_para(paragraphs[1],
             "AURIS: Çoklu-Model Topluluk Yaklaşımı ile Yapay Zekâ "
             "Tarafından Üretilen Müziklerin Tespiti",
             size=14, bold=True)

# TR yazar
replace_para(paragraphs[4], "Hasan Arthur Altuntaş1,*", size=9)

# TR kurum
replace_para(paragraphs[5],
             "1Düzce Üniversitesi, Mühendislik Fakültesi, "
             "Bilgisayar Mühendisliği Bölümü, 81620, Düzce, Türkiye",
             size=9)

# Boş 2. kurum satırı
replace_para(paragraphs[6], "", size=9)

# ORCID
replace_para(paragraphs[7], "0009-0002-8302-7657", size=9)

# E-posta
replace_para(paragraphs[8], "hasannarthurrr@gmail.com", size=9)

# Telefon (şablonun "İletişim yazarı telefon no:" altına ekle)
# P9 etiket, P10/P11 boş. P10'a telefon (varsa) — biz boş bırakıyoruz, sadece etiket
# Kullanıcı istemediği için boş

# EN başlık
replace_para(paragraphs[13],
             "AURIS: A Multi-Model Ensemble Approach for the Detection of "
             "AI-Generated Music",
             size=14, bold=True)

# EN yazar
replace_para(paragraphs[17], "Hasan Arthur Altuntaş1,*", size=9)

# EN kurum
replace_para(paragraphs[18],
             "1Department of Computer Engineering, Faculty of Engineering, "
             "Düzce University, 81620, Düzce, Türkiye",
             size=9)

# Boş 2. kurum EN
replace_para(paragraphs[19], "", size=9)

# ORCID EN
replace_para(paragraphs[20], "0009-0002-8302-7657", size=9)

# E-posta EN
replace_para(paragraphs[21], "hasannarthurrr@gmail.com", size=9)

doc.save(str(OUT))
print(f"Saved: {OUT}")
