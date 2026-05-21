"""Kontrol Listesi Formu — resmi şablonun sol sütun onay kutucuklarını tikler."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(r"D:/Downloads/Kontrol Listesi Formu.docx")
OUT = Path(__file__).resolve().parent / "deliverables" / "AURIS_Kontrol_Listesi_Formu.docx"


def _set_font(run, name="Times New Roman"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


doc = Document(str(SRC))

tbl = doc.tables[1]
for row in tbl.rows:
    cell = row.cells[0]
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    p = cell.add_paragraph()
    p.alignment = 1
    r = p.add_run("✓")
    _set_font(r)
    r.font.size = Pt(14)
    r.font.bold = True

doc.save(str(OUT))
print(f"Saved: {OUT}")
