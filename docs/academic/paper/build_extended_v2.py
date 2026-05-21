"""
Genişletilmiş İngilizce Özet — resmi şablonun (Genişletilmiş İngilizce Özet.docx)
üzerine birebir yaz.
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent
SRC = Path(r"D:/Downloads/Genişletilmiş İngilizce Özet.docx")
OUT = HERE / "deliverables" / "AURIS_Genisletilmis_Ingilizce_Ozet.docx"
FIGURES = HERE.parent / "figures"


def _set_font_basic(run, name="Times New Roman"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def cell_replace(cell, text, *, size=9, bold=False, italic=False):
    """Hücredeki tüm paragraf+run'ları temizle, tek paragraf+run ekle. Hücre stili korunur."""
    # Tüm paragrafları sil
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    # Yeni paragraf ekle
    p = cell.add_paragraph()
    r = p.add_run(text)
    _set_font_basic(r)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic


def cell_multi(cell, lines, *, size=9, bold=False):
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    for line in lines:
        p = cell.add_paragraph()
        r = p.add_run(line)
        _set_font_basic(r)
        r.font.size = Pt(size)
        r.font.bold = bold


def cell_image(cell, img_path, width_cm=6.5, caption=None):
    """Hücreyi temizle, görsel + altyazı ekle."""
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if img_path.exists():
        r = p.add_run()
        r.add_picture(str(img_path), width=Cm(width_cm))
    if caption:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(caption)
        _set_font_basic(r2)
        r2.font.size = Pt(8)
        r2.font.italic = True
        r2.font.bold = True


def main():
    doc = Document(str(SRC))

    # ===== Tablo 0: Başlık =====
    tbl0 = doc.tables[0]
    # R1C0 = "Title of the manuscript" → bizim EN başlığı
    cell_replace(tbl0.rows[1].cells[0],
                 "AURIS: A multi-model ensemble approach for the detection of "
                 "AI-generated music",
                 size=14, bold=True)
    # R2C0 = "ÖNEMLİ: ..." uyarısı — silebiliriz veya bırakabiliriz
    # Bırakıyoruz - kullanıcı submit ederken sileceğini biliyor

    # ===== Tablo 1: Asıl içerik (18 satır × 3 sütun) =====
    tbl = doc.tables[1]

    # R0: Yazar — 3 hücre aynı (template böyle)
    for c in range(3):
        cell_replace(tbl.rows[0].cells[c],
                     "Hasan Arthur Altuntaş1,*",
                     size=10, bold=True)

    # R1: Kurum (3 hücre aynı)
    for c in range(3):
        cell_replace(tbl.rows[1].cells[c],
                     "1Department of Computer Engineering, Faculty of "
                     "Engineering, Düzce University, 81620, Düzce, Türkiye",
                     size=9, italic=True)

    # R3C0: Highlights başlığı / R3C2: Graphical/Tabular Abstract başlığı
    cell_replace(tbl.rows[3].cells[0], "Highlights:", size=9, bold=True)
    cell_replace(tbl.rows[3].cells[2], "Graphical/Tabular Abstract",
                 size=9, bold=True)

    # R4C0: Highlights bullet'ları / R4C2: GA görseli (pipeline)
    cell_multi(tbl.rows[4].cells[0], [
        "• Proposed AURIS: end-to-end AI music detection with "
        "47-dimensional acoustic feature vector",
        "• 11-model ensemble (7 ML + 4 DL), LightGBM 95.48% ROC-AUC, "
        "lowest cross-fold variance (±0.0023)",
        "• Spectral flatness identified as most informative feature; "
        "17 of 47 features add no measurable accuracy",
    ], size=9)
    cell_image(tbl.rows[4].cells[2], FIGURES / "paper_pipeline_diagram.png",
               width_cm=6.5,
               caption="Figure A. AURIS system pipeline overview")

    # R5C0: Keywords başlığı
    cell_replace(tbl.rows[5].cells[0], "Keywords:", size=9, bold=True)
    cell_replace(tbl.rows[5].cells[2], "Purpose:", size=9, bold=True)

    # R6C0: Keywords / R6C2: Purpose text
    cell_multi(tbl.rows[6].cells[0], [
        "AI-generated music",
        "Deep learning",
        "Gradient boosting",
        "Ensemble learning",
        "Spectral flatness",
    ], size=9)
    cell_replace(tbl.rows[6].cells[2],
                 "The proliferation of text-to-music generators such as "
                 "Suno, Udio and MusicGen has raised a detection problem "
                 "of practical importance for copyright attribution, "
                 "streaming-platform integrity and artist economics. The "
                 "present study aims to design an end-to-end system that "
                 "distinguishes AI-generated music from human composition "
                 "with both competitive accuracy and operational "
                 "interpretability, while remaining deployable on "
                 "consumer-grade hardware.",
                 size=9)

    # R7C0: Article Info / R7C2: Theory and Methods başlık
    cell_replace(tbl.rows[7].cells[0], "Article Info:", size=9, bold=True)
    cell_replace(tbl.rows[7].cells[2], "Theory and Methods:", size=9, bold=True)

    # R8C0: Received/Accepted/DOI / R8C2: Theory text
    cell_multi(tbl.rows[8].cells[0], [
        "Research Article",
        "Received: dd.mm.yyyy",
        "Accepted: dd.mm.yyyy",
        "",
        "DOI:",
    ], size=9)
    cell_replace(tbl.rows[8].cells[2],
                 "A 47-dimensional acoustic feature vector covering "
                 "spectral, temporal, harmonic-tonal, MFCC and vocal "
                 "families is extracted using librosa. An ensemble of "
                 "eleven classifiers—Logistic Regression, Random Forest, "
                 "Gradient Boosting, SVM-RBF, MLP Neural Network, XGBoost, "
                 "LightGBM, Deep MLP, 1D-CNN, Residual MLP and Attention "
                 "MLP—is trained on 5,195 samples drawn from twelve AI "
                 "generation systems (Suno, Udio, MusicGen, AudioLDM2, "
                 "Stable Audio, Riffusion, Mustango, JEN-1, and others) "
                 "and multiple human sources (GTZAN, FMA, SleepyJesse "
                 "covers) under a stratified 5-fold cross-validation "
                 "protocol. Youden-J statistic is used for threshold "
                 "optimisation; Brier score and SHAP analysis for "
                 "calibration and interpretability respectively.",
                 size=9)

    # R9C2: Results başlık
    cell_replace(tbl.rows[9].cells[0], "", size=9)
    cell_replace(tbl.rows[9].cells[2], "Results:", size=9, bold=True)

    # R10C2: Results text
    cell_replace(tbl.rows[10].cells[0], "", size=9)
    cell_replace(tbl.rows[10].cells[2],
                 "LightGBM achieves the highest mean ROC-AUC at 0.9548 "
                 "with the lowest fold-to-fold standard deviation "
                 "(±0.0023) in the entire pool; its five folds span only "
                 "0.9515-0.9580. Deep MLP follows narrowly at 0.9542. "
                 "Spectral-flatness standard deviation ranks first in "
                 "feature importance. The Youden-optimal threshold "
                 "θ* = 0.4316 yields a balanced confusion matrix "
                 "(TN=2721, FP=392, FN=220, TP=1862) with 89.4% AI "
                 "sensitivity and 87.4% human specificity. A Brier score "
                 "of 0.083 confirms well-calibrated probabilities. "
                 "Per-source analysis reveals 93.0% recall on Suno and "
                 "88.6% on Echoes, but only 50.0% on the deepfake subset, "
                 "exposing the cross-generator generalisation challenge.",
                 size=9)

    # R11C0: Acknowledgement / R11C2: Conclusion başlık
    cell_replace(tbl.rows[11].cells[0], "Acknowledgement:", size=9, bold=True)
    cell_replace(tbl.rows[11].cells[2], "Conclusion:", size=9, bold=True)

    # R12C2: Conclusion text
    cell_replace(tbl.rows[12].cells[0], "—", size=9)
    cell_replace(tbl.rows[12].cells[2],
                 "The proposed AURIS system demonstrates that a "
                 "handcrafted 47-dimensional feature vector combined "
                 "with a heterogeneous ensemble of classifiers is "
                 "competitive with deep-learning alternatives for "
                 "AI-generated music detection. Diagnostic analysis "
                 "exposes two key limitations: tree-based models exhibit "
                 "an 8-14 percentage-point train-CV accuracy gap, and "
                 "the bottom 17 features add no measurable accuracy, "
                 "indicating clear opportunities for feature pruning and "
                 "overfitting mitigation. Future work will pursue formal "
                 "cross-generator held-out evaluation on emerging "
                 "benchmarks (SONICS, FakeMusicCaps), integration of "
                 "fine-tuned wav2vec2 into the protocol, explicit "
                 "adversarial-robustness testing under MP3 compression "
                 "and pitch shifting, and expansion of the dataset "
                 "toward ten thousand samples covering emerging "
                 "generation systems.",
                 size=9)

    # R13: boş
    cell_replace(tbl.rows[13].cells[0], "", size=9)
    cell_replace(tbl.rows[13].cells[2], "", size=9)

    # R14: Correspondence
    cell_replace(tbl.rows[14].cells[0], "Correspondence:", size=9, bold=True)
    cell_replace(tbl.rows[14].cells[2], "", size=9)

    # R15: Author/email/ORCID
    cell_multi(tbl.rows[15].cells[0], [
        "Author: Hasan Arthur Altuntaş",
        "e-mail: hasannarthurrr@gmail.com",
        "ORCID: 0009-0002-8302-7657",
    ], size=9)
    cell_replace(tbl.rows[15].cells[2], "", size=9)

    # R16: boş
    cell_replace(tbl.rows[16].cells[0], "", size=9)
    cell_replace(tbl.rows[16].cells[2], "", size=9)

    # R17: YAZARA NOT — şablonda hazır, dokunma

    doc.save(str(OUT))
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
