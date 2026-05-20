"""Kontrol Listesi Formu — sol sütundaki onay kutucuklarını işaretler."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(r"D:/Downloads/Kontrol Listesi Formu.docx")
OUT = Path(__file__).resolve().parent / "deliverables" / "AURIS_Kontrol_Listesi_Formu.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)


def _set_font(run, *, size=14, bold=True):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "Times New Roman")


def tick(cell):
    """Yazılı '✓' ekle (Word'deki form checkbox yerine basit görsel tik)."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = 1   # center
    r = p.add_run("✓")
    _set_font(r, size=14, bold=True)


doc = Document(str(SRC))

# Tablo 1: 17 satır. Tüm 17 maddeyi onayla
tbl = doc.tables[1]
for row in tbl.rows:
    tick(row.cells[0])

doc.save(str(OUT))
print(f"Saved: {OUT}")
