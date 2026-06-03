"""
Derin denetim:
1) Görsel en-boy oranı bozulması (aspect ratio distortion)
2) Metin sığmama tahmini (karakter yoğunluğu)
3) Şekil örtüşmesi (önemli metin kutuları üst üste mi)
"""
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

EMU = 914400
def inch(v): return v / EMU

base = 'docs/academic/TeslimEdilecekler/Sunumlar'
FIGS = 'docs/academic/figures'
files = ['01_Donem_Baslangic_Sunumu','02_Literatur_Taramasi_Sunumu',
         '03_Tasarim_ve_Gelistirme_Sunumu','04_Sonuc_ve_Demo_Sunumu',
         '05_Final_Savunma_Sunumu']

# ── 1) GÖRSEL ASPECT RATIO ──
print("=== GÖRSEL EN-BOY ORANI DENETİMİ ===")
# Slaytlardaki görsellerin yerleşim oranı vs orijinal oranı
img_placements = {
    'paper_pipeline_diagram.png': [(8.4,4.7),(8.4,4.7),(8.3,4.85)],
    'feature_distribution_ai_vs_human.png': [(8.3,4.85)],
    'paper_model_comparison.png': [(8.3,4.85)],
    'paper_score_distribution.png': [(4.8,3.9)],
    'paper_roc_curves.png': [(5.85,4.5)],
    'paper_confusion_matrix_lightgbm.png': [(5.0,4.5)],
    'shap_summary.png': [(6.0,4.85)],
    'threshold_sweep.png': [(6.15,3.05)],
}
for fname, placements in img_placements.items():
    path = f'{FIGS}/{fname}'
    if not os.path.exists(path):
        print(f"  EKSİK: {fname}")
        continue
    with Image.open(path) as im:
        ow, oh = im.size
    orig_ar = ow/oh
    for (pw, ph) in placements:
        place_ar = pw/ph
        distort = abs(orig_ar - place_ar) / orig_ar * 100
        flag = "  ⚠ BOZULMA" if distort > 8 else "  ✓"
        print(f"{flag} {fname[:38]:38s} orijinal {orig_ar:.2f} → yerleşim {place_ar:.2f}  ({distort:.0f}% sapma)")

# ── 2) METİN SIĞMAMA TAHMİNİ ──
print("\n=== METİN SIĞMAMA TAHMİNİ (yoğun kutular) ===")
issues = 0
for f in files:
    prs = Presentation(f'{base}/{f}.pptx')
    for si, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if not sh.has_text_frame: continue
            txt = sh.text_frame.text.strip()
            if not txt: continue
            try:
                w = inch(sh.width); h = inch(sh.height)
            except: continue
            # En büyük font boyutunu bul
            max_fs = 0
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        max_fs = max(max_fs, r.font.size.pt)
            if max_fs == 0: max_fs = 18
            # Kaba kapasite: kutu alanı (pt²) / (karakter genişlik×yükseklik)
            box_w_pt = w * 72; box_h_pt = h * 72
            char_w = max_fs * 0.52  # ortalama karakter genişliği
            line_h = max_fs * 1.25
            chars_per_line = max(1, box_w_pt / char_w)
            n_lines_needed = 0
            for line in txt.split('\n'):
                n_lines_needed += max(1, -(-len(line)//int(chars_per_line)))  # ceil
            lines_avail = box_h_pt / line_h
            if n_lines_needed > lines_avail * 1.15:  # %15 tolerans
                snip = (txt[:45]+'…') if len(txt)>45 else txt
                print(f"  ⚠ {f[:24]:24s} S{si+1}: ~{n_lines_needed} satır gerekli, "
                      f"~{lines_avail:.0f} sığar  fs={max_fs:.0f}  '{snip}'")
                issues += 1
print(f"\n  Potansiyel sığmama: {issues}")
