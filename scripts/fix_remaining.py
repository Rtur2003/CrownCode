import docx
from docx.shared import Pt
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')
paras = doc.paragraphs

print('Para 115:', paras[115].text)
print('Para 128:', paras[128].text)
print()

# Para 115: "Şekil 7" referansı
# "Öznitelik önemi analizi...Tablo 5 ilk on özniteliği, Şekil 7 önemi görselleştirmektedir"
# Şekil 6 = ilk yirmi öznitelik önem skorları (caption para 123) -> doğru atıf Şekil 6
# Şekil 7 = TreeSHAP -> ayrı atıf
# Para 115 büyük ihtimalle "Şekil 6" demeli öznitelik önem skorları için
p115 = paras[115]
t115 = p115.text
print('Para 115 full:', t115)
# Replace Şekil 7 with Şekil 6 in feature importance context
# But also check if TreeSHAP (Şekil 7) is mentioned separately
if 'Şekil 7' in t115 and 'TreeSHAP' not in t115:
    new115 = t115.replace('Şekil 7', 'Şekil 6')
    p115.clear()
    r = p115.add_run(new115)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9)
    print('Para 115: Şekil 7 -> Şekil 6')
elif 'Şekil 7' in t115:
    print('Para 115: Şekil 7 with TreeSHAP context - keeping as is')
else:
    print('Para 115: no Şekil 7 found')

# Para 128: "Şekil 10, 11, 12" check
# Captions: Şekil 10=Kalibrasyon, Şekil 11=Kesinlik-duyarlılık, Şekil 12=Karar eşiği
# Para 128 text mentions confusion matrix (Şekil 8) - check if 10,11,12 are forward refs
p128 = paras[128]
t128 = p128.text
print('\nPara 128 full:', t128)
# These are forward references to later figures - they should be correct as-is
# Şekil 10=Kalibrasyon, Şekil 11=PR curve, Şekil 12=threshold sweep - all valid

# Para 161: Şekil 17 -> Şekil 16
p161 = paras[161]
t161 = p161.text
print('\nPara 161 full:', t161)
if 'Şekil 17' in t161:
    new161 = t161.replace('Şekil 17', 'Şekil 16')
    p161.clear()
    r = p161.add_run(new161)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9)
    print('Para 161: Şekil 17 -> Şekil 16')

# Also check para 128 Şekil references are correct forward refs
# Şekil 8 (confusion matrix OK), Şekil 10 (calib), Şekil 11 (PR), Şekil 12 (threshold)
# These captions match: Şekil 10=Kalibrasyon ✓, Şekil 11=PR eğrisi ✓, Şekil 12=Eşik taraması ✓

doc.save('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')
print('\nKaydedildi.')
