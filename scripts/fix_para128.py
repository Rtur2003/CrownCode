import docx
from docx.shared import Pt
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')
paras = doc.paragraphs

# Captions ground truth:
# Şekil 8  = LightGBM karmaşıklık matrisi
# Şekil 9  = P(YZ) tahmin olasılık dağılımı
# Şekil 10 = Kalibrasyon eğrisi (Brier = 0,083)
# Şekil 11 = Kesinlik-duyarlılık eğrisi
# Şekil 12 = Karar eşiği taraması

# Para 128 currently says:
# "Şekil 10 insan ve YZ olasılık dağılımlarını..." -> should be Şekil 9
# "Şekil 11'deki kalibrasyon eğrisi..."           -> should be Şekil 10
# "Şekil 12 kesinlik-duyarlılık eğrisini..."      -> should be Şekil 11

p128 = paras[128]
t128 = p128.text
print('BEFORE:', t128)

# Replace in reverse order to avoid cascade
new128 = t128
new128 = new128.replace('Şekil 12 kesinlik', 'Şekil 11 kesinlik')
new128 = new128.replace("Şekil 11'deki kalibrasyon", "Şekil 10'daki kalibrasyon")
new128 = new128.replace('Şekil 10 insan', 'Şekil 9 insan')

if new128 != t128:
    p128.clear()
    r = p128.add_run(new128)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(9)
    print('AFTER:', new128)
else:
    print('No change')

doc.save('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')
print('Kaydedildi.')
