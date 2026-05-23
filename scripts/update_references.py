import docx
from docx.shared import Pt
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')


def set_para(p, text, size_pt=9):
    p.clear()
    run = p.add_run(text)
    run.bold = False
    run.font.size = Pt(size_pt)
    run.font.name = 'Times New Roman'


# ── Find reference paragraph indices ─────────────────────────────────────────
ref_map = {}
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    for n in range(1, 31):
        if t.startswith(f'{n}.') and len(t) > 10:
            ref_map[n] = i
            break

print('Bulunan referans paragrafları:', {k: v for k, v in sorted(ref_map.items())})

# ── Replacements ──────────────────────────────────────────────────────────────

# [5] preprint → Cros Vila et al. TISMIR 2025 (directly on AI music detection)
new_5 = ('5. Cros Vila L., Sturm B.L.T., Casini L., Dalmazzo D., '
         'The AI Music Arms Race: On the Detection of AI-Generated Music, '
         'Trans. Int. Soc. Music Inf. Retr., 8 (1), 179-194, 2025.')

# [6] preprint survey → Zhang et al. Sensors 2025 (peer-reviewed survey)
new_6 = ('6. Zhang B., Cui H., Nguyen V., Whitty M., '
         'Audio Deepfake Detection: What Has Been Achieved and What Lies Ahead, '
         'Sensors, 25 (7), 1989, 2025.')

# [8] preprint → SingFake ICASSP 2024 (published, singing voice deepfake detection)
new_8 = ('8. Zang Y., Zhang Y., Heydari M., Duan Z., '
         'SingFake: Singing Voice Deepfake Detection, '
         'Proceedings of the IEEE International Conference on Acoustics, Speech and Signal Processing '
         '(ICASSP 2024), 12156-12160, Seoul, Güney Kore, 14-19 Nisan, 2024.')

# [16] preprint → Xie et al. IEEE TIFS 2024 (published journal)
new_16 = ('16. Xie Y., Cheng H., Wang Y., Ye L., '
          'Domain Generalization via Aggregation and Separation for Audio Deepfake Detection, '
          'IEEE Trans. Inf. Forensics Secur., 19, 344-358, 2024.')

# [24] update arXiv → Journal of Imaging published version
new_24 = ('24. Comanducci L., Bestagini P., Tubaro S., '
          'FakeMusicCaps: A Dataset for Detection and Attribution of Synthetic Music Generated '
          'via Text-to-Music Models, J. Imaging, 11 (7), 242, 2025.')

# [25] preprint Echoes → Guo et al. ICASSP 2024 (self-supervised + multi-fusion for audio deepfake)
new_25 = ('25. Guo Y., Huang H., Chen X., Zhao H., Wang Y., '
          'Audio Deepfake Detection with Self-Supervised WavLM and Multi-Fusion Attentive Classifier, '
          'Proceedings of the IEEE International Conference on Acoustics, Speech and Signal Processing '
          '(ICASSP 2024), 12702-12706, Seoul, Güney Kore, 14-19 Nisan, 2024.')

# [26] coursework preprint → Yang et al. IEEE TIFS 2023 (multimodal deepfake, published journal)
new_26 = ('26. Yang W., Zhou X., Chen Z., Guo B., Ba Z., Xia Z., Cao X., Ren K., '
          'AVoiD-DF: Audio-Visual Joint Learning for Detecting Deepfake, '
          'IEEE Trans. Inf. Forensics Secur., 18, 2015-2029, 2023.')

# [27] LBD-only → Ding et al. Sensors 2024 (CNN music genre classification, published)
new_27 = ('27. Ding Y., Zhang H., Huang W., Zhou X., Shi Z., '
          'Efficient Music Genre Recognition Using ECAS-CNN: A Novel Channel-Aware Neural Network Architecture, '
          'Sensors, 24 (21), 7021, 2024.')

replacements = {
    5: new_5,
    6: new_6,
    8: new_8,
    16: new_16,
    24: new_24,
    25: new_25,
    26: new_26,
    27: new_27,
}

for num, new_text in replacements.items():
    idx = ref_map.get(num)
    if idx is not None:
        set_para(doc.paragraphs[idx], new_text, size_pt=9)
        print(f'[{num}] guncellendi.')
    else:
        print(f'[{num}] BULUNAMADI!')

doc.save('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')
print('Kaydedildi.')
