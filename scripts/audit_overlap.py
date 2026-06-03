"""Şekil örtüşme denetimi — metin kutuları üst üste mi?"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
EMU = 914400
def inch(v): return v / EMU

base = 'docs/academic/TeslimEdilecekler/Sunumlar'
files = ['05_Final_Savunma_Sunumu','03_Tasarim_ve_Gelistirme_Sunumu',
         '01_Donem_Baslangic_Sunumu','02_Literatur_Taramasi_Sunumu',
         '04_Sonuc_ve_Demo_Sunumu']

def overlap_area(a, b):
    l = max(a[0], b[0]); t = max(a[1], b[1])
    r = min(a[2], b[2]); btm = min(a[3], b[3])
    if r <= l or btm <= t: return 0
    return (r-l)*(btm-t)

total = 0
for f in files:
    prs = Presentation(f'{base}/{f}.pptx')
    for si, slide in enumerate(prs.slides):
        texts = []  # only text-bearing boxes
        for sh in slide.shapes:
            if not sh.has_text_frame: continue
            txt = sh.text_frame.text.strip()
            if not txt: continue
            try:
                l=inch(sh.left); t=inch(sh.top)
                r=l+inch(sh.width); b=t+inch(sh.height)
            except: continue
            texts.append((l,t,r,b,txt))
        for i in range(len(texts)):
            for j in range(i+1, len(texts)):
                a, b = texts[i], texts[j]
                ov = overlap_area(a, b)
                area_a = (a[2]-a[0])*(a[3]-a[1])
                area_b = (b[2]-b[0])*(b[3]-b[1])
                small = min(area_a, area_b)
                if small > 0 and ov/small > 0.35:  # %35+ örtüşme
                    sa = (a[4][:30]+'…') if len(a[4])>30 else a[4]
                    sb = (b[4][:30]+'…') if len(b[4])>30 else b[4]
                    print(f"  {f[:22]:22s} S{si+1}: %{ov/small*100:.0f} örtüşme")
                    print(f"      A: '{sa}'  [{a[0]:.2f},{a[1]:.2f},{a[2]:.2f},{a[3]:.2f}]")
                    print(f"      B: '{sb}'  [{b[0]:.2f},{b[1]:.2f},{b[2]:.2f},{b[3]:.2f}]")
                    total += 1
print(f"\nTOPLAM ÖRTÜŞME: {total}")
