import docx
from docx.shared import Pt
import re
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')

def set_para(p, text, size_pt=9):
    p.clear()
    run = p.add_run(text)
    run.bold = False
    run.font.size = Pt(size_pt)
    run.font.name = 'Times New Roman'

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if re.match(r'^29\.', t) and 'GTZAN' in t:
        print(f'Bulundu para {i}: {t}')
        new_text = ('29. Özbalcı M.C., Şahin H., Bilgin T.T., '
                    'Classification of Music Genres of GTZAN Dataset with Machine Learning Methods, '
                    'Mühendislik Bilimleri ve Araştırmaları Dergisi (J. Eng. Sci. Res.), 6 (1), 77-87, 2024.')
        set_para(p, new_text, size_pt=9)
        print(f'Guncellendi: {new_text}')
        break
else:
    print('Ref [29] bulunamadi!')

doc.save('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')
print('Kaydedildi.')
