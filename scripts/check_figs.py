import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('docs/academic/paper/deliverables/AURIS_Makale_Metni.docx')

print('=== ALL FIGURE CAPTIONS ===')
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith('Şekil ') and '.' in t[:10]:
        print(f'  Para {i}: {t[:100]}')

print('\n=== ALL FIGURE REFERENCES IN BODY ===')
import re
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    refs = re.findall(r'Şekil\s+\d+', t)
    # exclude captions
    if refs and not (t.startswith('Şekil ') and '.' in t[:10]):
        print(f'  Para {i}: {refs} -> {t[:120]}')
